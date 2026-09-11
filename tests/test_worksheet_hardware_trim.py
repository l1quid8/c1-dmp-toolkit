"""Export only actual 505/714 hardware, with complete blocks and closed borders."""

import sys
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from parse_dmp_worksheet import DMPDesign, PowerSupply, RSP, parse_dmp_worksheet

TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
PS_SHEET = "DMP 505-12_G Power Supply 1-10"
RSP_SHEET = "DMP 714 Exp Mod"
pytestmark = pytest.mark.skipif(not TEMPLATE.exists(), reason="Worksheet template unavailable")


def assert_table_ends_at(ws, last_row):
    assert ws.max_row == last_row, "unused template rows remain"
    assert max(ws.row_dimensions, default=1) <= last_row
    assert all(merge.max_row <= last_row for merge in ws.merged_cells)
    assert range_boundaries(ws.print_area.rsplit("!", 1)[-1]) == (1, 1, 4, last_row)
    for col in range(1, 5):
        assert ws.cell(last_row, col).border.bottom.style == "medium"
    assert ws.cell(last_row, 1).border.left.style == "thick"
    assert ws.cell(last_row, 4).border.right.style == "thick"


@pytest.mark.parametrize("count,last_row", [(0, 1), (1, 5), (7, 29), (30, 121)])
def test_power_supply_sheet_keeps_only_complete_installed_blocks(tmp_path, count, last_row):
    design = DMPDesign(power_supplies=[
        PowerSupply(n, f"PS ROOM {n}", {2: f"AC {n}", 3: f"BATT {n}"})
        for n in range(1, count + 1)
    ])
    output = tmp_path / "power-supplies.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        ws = wb[PS_SHEET]
        assert_table_ends_at(ws, last_row)
        assert ws["A1"].value == "DMP505-12G\nNUMBER"
        assert "B1:C1" in ws.merged_cells
        for i, ps in enumerate(design.power_supplies):
            row = 2 + 4 * i
            assert ws.cell(row, 1).value == ps.number
            assert f"A{row}:A{row + 3}" in ws.merged_cells
            assert [ws.cell(row + j, 2).value for j in range(4)] == [
                "RELAY 1", "RELAY 2", "RELAY 3", "RELAY 4"]
            assert [ws.cell(row + j, 3).value for j in range(4)] == [
                "12v DC Output to Terminal Strip", f"AC {ps.number}", f"BATT {ps.number}", None]
            assert [ws.cell(row + j, 4).value for j in range(4)] == [ps.location] * 4
        if count:
            assert ws.cell(last_row - 3, 1).border.bottom.style == "medium"
    finally:
        wb.close()


@pytest.mark.parametrize("count,last_row", [(0, 3), (1, 4), (7, 10), (15, 18)])
def test_expander_sheet_keeps_headers_and_each_installed_module(tmp_path, count, last_row):
    # Both 8- and 16-point modules must retain their actual model and zone range.
    design = DMPDesign(rsps=[
        RSP(n, f"RSP ROOM {n}", list(range(501 + 16 * (n - 1),
                                         501 + 16 * (n - 1) + (8 if n % 2 else 16))))
        for n in range(1, count + 1)
    ])
    output = tmp_path / "expanders.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        ws = wb[RSP_SHEET]
        assert_table_ends_at(ws, last_row)
        assert ws["A1"].value == "DMP 714-16 or 714-08 EXPANSION MODULES"
        assert set(map(str, ws.merged_cells)) == {"A1:D2", "A3:B3"}
        for i, rsp in enumerate(design.rsps):
            row = 4 + i
            assert ws.cell(row, 1).value == ("DMP 714-8 #" if rsp.number % 2 else "DMP 714-16 #")
            assert ws.cell(row, 2).value == rsp.number
            assert ws.cell(row, 3).value == f"{min(rsp.zones)} - {max(rsp.zones)}"
            assert ws.cell(row, 4).value == rsp.location
    finally:
        wb.close()


def test_sparse_device_ids_do_not_create_placeholders_or_lose_relay_data(tmp_path):
    design = DMPDesign(
        rsps=[RSP(2, "OFFICE", list(range(501, 509))),
              RSP(7, "CAFETERIA", list(range(509, 525)))],
        power_supplies=[PowerSupply(2, "OFFICE", {4: "Spare"}),
                        PowerSupply(7, "CAFETERIA", {4: "Spare"})],
    )
    output = tmp_path / "sparse.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        assert_table_ends_at(wb[PS_SHEET], 9)
        assert_table_ends_at(wb[RSP_SHEET], 5)
    finally:
        wb.close()
    reread = parse_dmp_worksheet(output)
    assert reread.rsps == design.rsps
    assert [ps.number for ps in reread.power_supplies] == [2, 7]
    for actual, expected in zip(reread.power_supplies, design.power_supplies):
        assert actual.location == expected.location
        for relay in (2, 3, 4):
            assert actual.relays[relay] == expected.relays[relay]


@pytest.mark.parametrize("count", [8, 15, 30])
def test_power_supply_import_keeps_every_block_and_its_final_relays(tmp_path, count):
    design = DMPDesign(power_supplies=[
        PowerSupply(n, f"PS ROOM {n}", {2: f"AC {n}", 3: f"BATT {n}", 4: "Spare"})
        for n in range(1, count + 1)
    ])
    output = tmp_path / "power-supply-roundtrip.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    actual = parse_dmp_worksheet(output).power_supplies
    assert [ps.number for ps in actual] == list(range(1, count + 1))
    for ps, expected in zip(actual, design.power_supplies):
        assert ps.location == expected.location
        assert ps.relays == expected.relays
