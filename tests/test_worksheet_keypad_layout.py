"""Keypad exports must extend the template's six-row table for every device."""

import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from parse_dmp_worksheet import DMPDesign, Keypad, parse_dmp_worksheet

TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
pytestmark = pytest.mark.skipif(not TEMPLATE.exists(), reason="Worksheet template unavailable")


@pytest.mark.parametrize("count", [0, 1, 6, 8, 28])
def test_keypad_table_formats_every_device_and_ends_at_actual_count(tmp_path, count):
    design = DMPDesign(keypads=[
        Keypad(n, "MSP" if n == 1 else "710-KP-1", f"KEYPAD ROOM {n}", n <= 6)
        for n in range(1, count + 1)
    ])
    output = tmp_path / "keypads.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    template = openpyxl.load_workbook(TEMPLATE)
    try:
        ws = wb["Keypad"]
        reference = template["Keypad"]
        last_row = 2 + count
        for row in range(3, last_row + 1):
            assert ws.row_dimensions[row].height == reference.row_dimensions[3].height
            for col in range(1, 5):
                cell = ws.cell(row, col)
                expected = reference.cell(3 if row == 3 else 4, col)
                assert cell.font == copy(expected.font), f"{cell.coordinate}: keypad font changed"
                assert cell.alignment == copy(expected.alignment)
                assert cell.fill == copy(expected.fill)
                assert cell.number_format == expected.number_format
                assert cell.border.left == expected.border.left
                assert cell.border.right == expected.border.right
                assert cell.border.top == expected.border.top
                assert cell.border.bottom.style == ("thick" if row == last_row else "hair")

        assert ws.max_row == last_row, "unused keypad rows remain"
        assert ws.max_column == 4
        assert max(ws.row_dimensions, default=1) <= last_row
        assert range_boundaries(ws.print_area.rsplit("!", 1)[-1]) == (1, 1, 4, last_row)
        assert ws.print_title_rows == "$1:$2"
        assert set(map(str, ws.merged_cells)) == {"A1:D1"}
        for row in (1, 2):
            for col in range(1, 5):
                assert ws.cell(row, col).value == reference.cell(row, col).value
    finally:
        template.close()
        wb.close()

    assert parse_dmp_worksheet(output).keypads == design.keypads


def test_sparse_keypad_numbers_keep_data_without_placeholder_rows(tmp_path):
    design = DMPDesign(keypads=[
        Keypad(1, "MSP", "ADMIN BUILDING (COORDINATOR) (Service Keypad)", True),
        Keypad(7, "710-KP-5", "EARLY EDUCATION CENTER (WAITING AREA)", False),
        Keypad(8, "710-KP-4", "CAFETERIA BLDG (KITCHEN)", False),
    ])
    output = tmp_path / "sparse-keypads.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["Keypad"]
        assert ws.max_row == 5
        assert [ws.cell(row, 1).value for row in range(3, 6)] == [1, 7, 8]
        assert all(ws.cell(5, col).border.bottom.style == "thick" for col in range(1, 5))
    finally:
        wb.close()
    assert parse_dmp_worksheet(output).keypads == design.keypads
