"""BOM imports preserve evidence and leave unknown locations for review."""

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from parse_bom import parse_bom, NotBOMError
from validation import validate_design


def test_bom_breakdown_builds_reviewable_design(tmp_path):
    path = tmp_path / "site.xlsx"
    wb = openpyxl.Workbook()
    equipment = wb.active
    equipment.title = "DMP Equipment"
    equipment["A1"] = "TEST SCHOOL"
    for row, model, quantity in ((3, "XR550", 1), (4, "PS12-5", 2),
                                 (5, "714-16", 1), (6, "714-8", 1),
                                 (7, "710", 2), (8, "7000 Series", 1),
                                 (9, "DS9370", 2)):
        equipment.cell(row, 2, model)
        equipment.cell(row, 4, quantity)
    bom = wb.create_sheet("BOM")
    bom["A1"] = "TEST SCHOOL"
    bom["A2"] = "123 TEST ST, TEST CITY, CA 90000, LOCATION ID: 42"
    bom["B4"] = "MANUFACTURER"
    breakdown = wb.create_sheet("BOM Breakdown")
    for row, number, start, count, location in (
        (1, 1, 501, 14, "ADMIN ROOM - BLDG. A"),
        (20, 2, 517, 6, "CLASSROOM 2 - BLDG. B"),
    ):
        breakdown.cell(row, 2, f"RSP {number}")
        breakdown.cell(row, 19, location)
        for offset in range(count):
            breakdown.cell(row, 3 + offset, f"Z{start + offset}")
        breakdown.cell(row + 14, 1, "Motion Sens")
    breakdown.cell(15, 3, 1)
    breakdown.cell(34, 4, 1)
    breakdown.cell(4, 19, "FROM KP-710-1 TO MAIN OFFICE KEYPAD")
    wb.save(path)

    result = parse_bom(path)

    assert result.design.site_info.school_name == "TEST SCHOOL"
    assert result.design.site_info.address_line1 == "123 TEST ST"
    assert [(r.number, r.model, r.location, len(r.zones)) for r in result.design.rsps] == [
        (1, "714-16", "ADMIN ROOM - BLDG. A", 16),
        (2, "714-8", "CLASSROOM 2 - BLDG. B", 8),
    ]
    assert [(z.number, z.device_type) for z in result.design.zones if z.device_type == "Motion"] == [
        (501, "Motion"), (518, "Motion")]
    assert all("NEEDS REVIEW" in z.location for z in result.design.zones
               if z.device_type == "Motion")
    assert len(result.design.splitters) == 1
    assert len(result.design.keypads) == 1
    assert "BOM Breakdown!S20" in result.review_text
    assert "BOM Breakdown!S4" in result.review_text
    assert any(issue.code == "zone.location_review" and issue.ref == "zone:501"
               for issue in validate_design(result.design, topology_confirmed=False))


def test_unrelated_workbook_is_not_bom(tmp_path):
    path = tmp_path / "other.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(NotBOMError):
        parse_bom(path)
