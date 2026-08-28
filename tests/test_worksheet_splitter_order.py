"""Sort whole splitter blocks before populating the merged-cell worksheet."""

import copy
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from parse_dmp_worksheet import DMPDesign, Splitter
from riser_model import DevicePortRef
from topology_service import connect, project_legacy_topology

TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
pytestmark = pytest.mark.skipif(not TEMPLATE.exists(), reason="Worksheet template unavailable")


def mixed_design():
    ids = ["710-LX900-2", "710-LX500-10", "710-KP-10", "710-LX600-2",
           "710-LX500-2", "710-KP-2", "710-LX800-1", "710-LX700-1",
           "710-LX500-1", "710-KP-1", "710-LX600-1", "710-LX900-1"]
    design = DMPDesign(splitters=[
        Splitter(sid, "KP" if "KP" in sid else "LX", f"ROOM {sid}",
                 outputs=["Spare"] * 3) for sid in ids])
    # Numeric order must not change electrical order: -10 feeds -2.
    connect(design, DevicePortRef("MSP", "LX500"), DevicePortRef("710-LX500-10", "IN"))
    connect(design, DevicePortRef("710-LX500-10", "OUT2"), DevicePortRef("710-LX500-2", "IN"))
    connect(design, DevicePortRef("MSP", "PROG"), DevicePortRef("710-KP-10", "IN"))
    connect(design, DevicePortRef("710-KP-10", "OUT3"), DevicePortRef("710-KP-2", "IN"))
    project_legacy_topology(design)
    return design


@pytest.fixture
def generated(tmp_path):
    design = mixed_design()
    before = copy.deepcopy(design)
    out = tmp_path / "ordered.xlsx"
    write_dmp_xlsx(design, TEMPLATE, out)
    wb = openpyxl.load_workbook(out)
    yield design, before, wb
    wb.close()


@pytest.mark.parametrize("sheet,expected", [
    ("710 Splitter-Repeater LX500", ["710-LX500-1", "710-LX500-2", "710-LX500-10",
                                    "710-LX600-1", "710-LX600-2", "710-LX700-1",
                                    "710-LX800-1", "710-LX900-1", "710-LX900-2"]),
    ("710 Splitter-Repeater(KP-Bus) ", ["710-KP-1", "710-KP-2", "710-KP-10"]),
])
def test_worksheet_contains_only_real_splitters_in_numeric_order(generated, sheet, expected):
    _, _, wb = generated
    ws = wb[sheet]
    actual = [ws.cell(row, 1).value for row in range(2, ws.max_row + 1)
              if ws.cell(row, 1).value]
    assert actual == expected


def test_sorted_blocks_keep_locations_ports_merges_and_topology(generated):
    design, before, wb = generated
    assert design.splitters == before.splitters
    assert design.connections == before.connections
    template = openpyxl.load_workbook(TEMPLATE)
    try:
        for family, sheet, original in (
                ("KP", "710 Splitter-Repeater(KP-Bus) ", "710 Splitter-Repeater(KP-Bus) "),
                ("LX", "710 Splitter-Repeater LX500", "710 Splitter-Repeater LXBus")):
            ws = wb[sheet]
            used_rows = 1 + 4 * sum(s.splitter_type == family for s in before.splitters)
            assert set(map(str, ws.merged_cells)) == {
                str(merge) for merge in template[original].merged_cells
                if merge.max_row <= used_rows}
            for splitter in (s for s in before.splitters if s.splitter_type == family):
                row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == splitter.id)
                assert (row - 2) % 4 == 0
                assert ws.cell(row, 4).value == splitter.location
                assert [ws.cell(row + i, 3).value for i in (1, 2, 3)] == splitter.outputs
                for i, target in enumerate(splitter.outputs, 1):
                    if target.startswith("To "):
                        assert ws.cell(row + i, 4).value == f"ROOM {target[3:]}"
                for input_value in splitter.inputs.values():
                    if input_value:
                        assert ws.cell(row, 3).value == input_value
    finally:
        template.close()


def test_empty_design_does_not_export_template_splitters(tmp_path):
    out = tmp_path / "empty.xlsx"
    write_dmp_xlsx(DMPDesign(), TEMPLATE, out)
    wb = openpyxl.load_workbook(out)
    try:
        for ws in wb:
            if "710" in ws.title:
                assert all(ws.cell(row, 1).value is None for row in range(2, ws.max_row + 1))
    finally:
        wb.close()


@pytest.mark.parametrize("sheet,last_row,last_output", [
    ("710 Splitter-Repeater LX500", 37, "Spare"),
    ("710 Splitter-Repeater(KP-Bus) ", 13, "To 710-KP-2"),
])
def test_710_sheets_end_after_last_real_splitter_block(generated, sheet, last_row, last_output):
    _, _, wb = generated
    ws = wb[sheet]

    assert ws.max_row == last_row, "blank template cells still extend the used range"
    assert max(ws.row_dimensions, default=1) <= last_row, "orphan row formatting remains"
    assert all(merge.max_row <= last_row for merge in ws.merged_cells)
    assert ws.print_area, "printing is not limited to real device blocks"
    assert range_boundaries(ws.print_area.rsplit("!", 1)[-1]) == (1, 1, 4, last_row)
    assert ws.cell(last_row, 2).value.endswith("Bus 3"), "do not trim real spare ports"
    assert ws.cell(last_row, 3).value == last_output


@pytest.mark.parametrize("sheet,last_row", [
    ("710 Splitter-Repeater LX500", 37),
    ("710 Splitter-Repeater(KP-Bus) ", 13),
])
def test_trimmed_710_table_has_a_continuous_bottom_border(generated, sheet, last_row):
    _, _, wb = generated
    ws = wb[sheet]
    for col in range(1, 5):
        assert ws.cell(last_row, col).border.bottom.style == "medium"
    # Excel takes the merged ID cell's border from the anchor as well as its edge.
    assert ws.cell(last_row - 3, 1).border.bottom.style == "medium"
    assert ws.cell(last_row, 1).border.left.style == "thick"
    assert ws.cell(last_row, 4).border.right.style == "thick"


@pytest.mark.parametrize("family,count,last_row", [
    ("LX", 0, 1), ("KP", 0, 1), ("LX", 1, 5), ("KP", 1, 5),
    ("LX", 20, 81), ("KP", 8, 33),
])
def test_trimming_handles_empty_single_and_full_template(tmp_path, family, count, last_row):
    prefix = "710-LX900" if family == "LX" else "710-KP"
    design = DMPDesign(splitters=[Splitter(f"{prefix}-{n}", family,
                                          outputs=["Spare"] * 3)
                                 for n in range(1, count + 1)])
    out = tmp_path / "trimmed.xlsx"
    write_dmp_xlsx(design, TEMPLATE, out)
    wb = openpyxl.load_workbook(out)
    try:
        ws = wb["710 Splitter-Repeater LX500" if family == "LX"
                else "710 Splitter-Repeater(KP-Bus) "]
        assert ws.max_row == last_row
        assert max(ws.row_dimensions, default=1) <= last_row
        assert ws.print_area, "printing is not limited to real device blocks"
        assert range_boundaries(ws.print_area.rsplit("!", 1)[-1]) == (1, 1, 4, last_row)
        assert ws["A1"].value == "710 Bus Splitter/Repeater Modules"
        assert all(ws.cell(last_row, col).border.bottom.style == "medium"
                   for col in range(1, 5))
        if count:
            assert ws.cell(last_row - 3, 1).value == f"{prefix}-{count}"
            assert ws.cell(last_row, 3).value == "Spare"
            assert f"A{last_row - 3}:A{last_row}" in ws.merged_cells
    finally:
        wb.close()
