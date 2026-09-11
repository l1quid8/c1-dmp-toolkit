"""Actual LX addresses outside the old 96-point template remain exportable."""

from copy import copy
from pathlib import Path
import sys
import zipfile

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from parse_dmp_worksheet import DMPDesign, RSP, ZoneInfo, parse_dmp_worksheet

TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
pytestmark = pytest.mark.skipif(not TEMPLATE.exists(), reason="Worksheet template unavailable")


def test_full_lx_address_bands_preserve_edge_points_and_descriptions(tmp_path):
    # Include both ends of every bus. These existing imported ranges must retain
    # their addresses even though new projects normally begin at x01.
    modules = [
        RSP(number, f"ROOM {number}", list(range(start, start + 8)), "714-8")
        for number, start in enumerate(
            (start for bus in range(500, 901, 100) for start in (bus, bus + 92)), 1
        )
    ]
    zones = [
        ZoneInfo(
            number,
            "SPARE" if index == 0 else
            f"PS-{module.number}: A/C LOSS" if index == 6 else
            f"PS-{module.number}: BATT. TRBL" if index == 7 else
            f"ROOM POINT {number}",
            "Spare" if index == 0 else "Supervisory" if index >= 6 else "Motion",
            1,
        )
        for module in modules for index, number in enumerate(module.zones)
    ]
    design = DMPDesign(rsps=modules, zones=zones)
    before = [(module.number, list(module.zones)) for module in modules]
    output = tmp_path / "full-bus-edge-addresses.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)

    assert [(module.number, module.zones) for module in design.rsps] == before
    wb = openpyxl.load_workbook(output)
    template = openpyxl.load_workbook(TEMPLATE)
    try:
        master = wb["Master"]
        by_number = {
            int(row[0].value[1:]): row[0].row
            for row in master.iter_rows(min_row=2, max_col=2)
            if isinstance(row[0].value, str) and row[0].value.startswith("Z")
        }
        assert {zone.number for zone in zones} <= by_number.keys(), "Master omitted valid LX addresses"
        # Existing Master references remain stable; extra occupied points extend
        # the styled table instead of shifting the template's formulas.
        for row in template["Master"].iter_rows(min_row=2, max_col=1):
            assert master.cell(row[0].row, 1).value == row[0].value
        assert master.max_row == template["Master"].max_row + 20
        table = master.tables["Sheet1"]
        assert table.ref == f"A1:B{master.max_row}"
        assert table.autoFilter.ref == table.ref
        for zone in zones:
            mrow = by_number[zone.number]
            assert master.cell(mrow, 2).value == zone.location
            if mrow > template["Master"].max_row:
                for column in (1, 2):
                    assert master.cell(mrow, column)._style == copy(template["Master"].cell(481, column)._style)
                assert any(f"B{mrow}" in formatting.sqref for formatting in master.conditional_formatting)

        current_bus = None
        summary = {}
        for row in wb["DMP XR550"].iter_rows(min_row=13, max_col=4, values_only=True):
            current_bus = row[0] or current_bus
            if row[1]:
                summary[int(row[1].rsplit("-", 1)[-1])] = (current_bus, row[2])
        for module in modules:
            bus = module.zones[0] // 100 * 100
            assert summary[module.number] == (f"LX Bus {bus}", f"{module.zones[0]}-{module.zones[-1]}")
            info = wb[f"DMP 714-16 Point Info ({module.number})"]
            for row, number in enumerate(module.zones, 4):
                mrow = by_number[number]
                assert info.cell(row, 1).value == f"=Master!A{mrow}"
                assert info.cell(row, 2).value == f"=Master!B{mrow}"
                assert info.cell(row, 1)._style == copy(template["DMP 714 Point Info 1"].cell(row, 1)._style)
    finally:
        template.close()
        wb.close()

    # Read-only consumers obey the stored used range; it must include the rows
    # appended after openpyxl restores the template's Excel-specific metadata.
    readonly = openpyxl.load_workbook(output, read_only=True)
    try:
        assert readonly["Master"].max_row == 501
    finally:
        readonly.close()
    with zipfile.ZipFile(TEMPLATE) as original, zipfile.ZipFile(output) as generated:
        assert generated.read("xl/tables/table1.xml") == original.read("xl/tables/table1.xml").replace(
            b'ref="A1:B481"', b'ref="A1:B501"'
        ), "Only the table bounds should change; retain external-query metadata"

    reread = parse_dmp_worksheet(output)
    descriptions = {zone.number: zone.description for zone in reread.master_zones}
    assert reread.rsps == design.rsps
    assert {zone.number: descriptions[zone.number] for zone in zones} == {
        zone.number: zone.location for zone in zones
    }


def test_master_adds_occupied_edge_zone_without_an_rsp(tmp_path):
    output = tmp_path / "unassigned-edge-zone.xlsx"
    write_dmp_xlsx(DMPDesign(zones=[ZoneInfo(599, "UNASSIGNED ROOM", "Motion")]), TEMPLATE, output)
    reread = parse_dmp_worksheet(output)
    assert any(zone.number == 599 and zone.description == "UNASSIGNED ROOM" for zone in reread.master_zones)
