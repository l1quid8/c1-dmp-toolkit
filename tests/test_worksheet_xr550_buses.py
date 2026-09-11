"""The XR550 summary must list every RSP under its actual zone-address bus."""

from copy import copy
from pathlib import Path
import sys

import openpyxl
from openpyxl.utils.cell import range_boundaries
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from hardware import MAX_EXPANDERS, MAX_KEYPADS, MAX_SPLITTERS_PER_TYPE, add_expander
from parse_dmp_worksheet import DMPDesign, Keypad, RSP, Splitter, parse_dmp_worksheet

TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
pytestmark = pytest.mark.skipif(not TEMPLATE.exists(), reason="Worksheet template unavailable")


def _summary(ws):
    bus = None
    result = []
    for row in ws.iter_rows(min_row=13, max_col=4, values_only=True):
        if row[0]:
            bus = row[0]
        if row[1]:
            result.append((bus, *row[1:]))
    return result


def test_seventh_packed_rsp_stays_on_bus_500_with_complete_table(tmp_path):
    design = DMPDesign(rsps=[
        RSP(n, f"ROOM {n}", list(range(start, end + 1)))
        for n, start, end in [(1, 501, 516), (2, 517, 524), (3, 525, 540),
                              (4, 541, 548), (5, 549, 564), (6, 565, 580), (7, 581, 596)]
    ])
    design.rsps[-1].location = "CAFETERIA BLDG. (DRY STORAGE)"
    output = tmp_path / "seven-rsps.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    template = openpyxl.load_workbook(TEMPLATE)
    try:
        ws = wb["DMP XR550"]
        rows = _summary(ws)
        assert len(rows) == 7, "XR550 summary dropped an installed RSP"
        assert rows[-1] == ("LX Bus 500", "714-16-7", "581-596", "CAFETERIA BLDG. (DRY STORAGE)")
        assert all(row[0] == "LX Bus 500" for row in rows)
        assert "A13:A19" in ws.merged_cells
        assert ws["A20"].value == "LX Bus 600"
        assert "A20:A25" in ws.merged_cells
        assert "D19:E19" in ws.merged_cells
        assert ws["B19"].font == copy(template["DMP XR550"]["B14"].font)
        assert ws["D19"].alignment == copy(template["DMP XR550"]["D14"].alignment)
        assert ws["A19"].border.left.style == "thick"
        assert ws["E19"].border.right.style == "thick"
        assert ws.row_dimensions[19].height == 25
        assert ws.max_row == 43
        assert range_boundaries(ws.print_area.rsplit("!", 1)[-1]) == (1, 1, 5, 43)
        assert all(ws.cell(43, col).border.bottom.style == "thick" for col in range(1, 6))
        for row in range(1, 13):
            for col in range(1, 6):
                assert ws.cell(row, col).value == template["DMP XR550"].cell(row, col).value
                assert ws.cell(row, col)._style == template["DMP XR550"].cell(row, col)._style
    finally:
        template.close()
        wb.close()


def test_twelve_eight_point_modules_expand_one_bus_without_overwriting_next(tmp_path):
    design = DMPDesign(rsps=[
        RSP(n, f"ROOM {n}", list(range(501 + 8 * (n - 1), 509 + 8 * (n - 1))))
        for n in range(1, 13)
    ] + [RSP(13, "BUS 600 ROOM", list(range(601, 617)))])
    output = tmp_path / "twelve-on-one-bus.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["DMP XR550"]
        rows = _summary(ws)
        assert len(rows) == 13
        assert rows[11] == ("LX Bus 500", "714-8-12", "589-596", "ROOM 12")
        assert rows[12] == ("LX Bus 600", "714-16-13", "601-616", "BUS 600 ROOM")
        assert "A13:A24" in ws.merged_cells
        assert ws["A25"].value == "LX Bus 600"
        assert "A25:A30" in ws.merged_cells
    finally:
        wb.close()


def test_bus_labels_follow_zone_addresses_for_sparse_out_of_order_ids(tmp_path):
    design = DMPDesign(rsps=[
        RSP(1, "BUS 900", list(range(901, 917))),
        RSP(15, "BUS 500", list(range(501, 509))),
        RSP(4, "BUS 800", list(range(801, 817))),
        RSP(11, "BUS 600", list(range(601, 617))),
        RSP(7, "BUS 700", list(range(701, 709))),
    ])
    output = tmp_path / "sparse-buses.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        assert _summary(wb["DMP XR550"]) == [
            ("LX Bus 500", "714-8-15", "501-508", "BUS 500"),
            ("LX Bus 600", "714-16-11", "601-616", "BUS 600"),
            ("LX Bus 700", "714-8-7", "701-708", "BUS 700"),
            ("LX Bus 800", "714-16-4", "801-816", "BUS 800"),
            ("LX Bus 900", "714-16-1", "901-916", "BUS 900"),
        ]
    finally:
        wb.close()


@pytest.mark.parametrize("zones,zone_range", [
    ([], None),
    (list(range(593, 609)), "593-608"),
    (list(range(492, 500)), "492-499"),
    (list(range(1000, 1008)), "1000-1007"),
])
def test_rsp_with_unknown_or_invalid_bus_addresses_remains_visible(tmp_path, zones, zone_range):
    output = tmp_path / "needs-review.xlsx"
    write_dmp_xlsx(DMPDesign(rsps=[RSP(7, "UNASSIGNED ROOM", zones)]), TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        assert _summary(wb["DMP XR550"]) == [
            ("Bus needs review", "714-8-7" if len(zones) == 8 else "714-16-7", zone_range, "UNASSIGNED ROOM")
        ]
    finally:
        wb.close()


def test_maximum_supported_inventory_survives_all_worksheet_hardware_sheets(tmp_path):
    design = DMPDesign()
    for number in range(1, MAX_EXPANDERS + 1):
        add_expander(design, "714-8" if number % 2 else "714-16", f"ROOM {number}")
    design.keypads = [Keypad(n, "MSP", f"KEYPAD {n}", True) for n in range(1, MAX_KEYPADS + 1)]
    design.splitters = [
        Splitter(f"710-{prefix}-{n}", kind, f"SPLITTER ROOM {n}")
        for prefix, kind in [("KP", "KP"), ("LX500", "LX")]
        for n in range(1, MAX_SPLITTERS_PER_TYPE + 1)
    ]
    output = tmp_path / "maximum-inventory.xlsx"
    write_dmp_xlsx(design, TEMPLATE, output)
    wb = openpyxl.load_workbook(output)
    try:
        summary = _summary(wb["DMP XR550"])
        assert len(summary) == MAX_EXPANDERS
        assert [row[0] for row in summary] == ["LX Bus 500"] * 8 + ["LX Bus 600"] * 7
        assert {int(row[1].rsplit("-", 1)[-1]) for row in summary} == set(range(1, MAX_EXPANDERS + 1))
        for rsp in design.rsps:
            info = wb[f"DMP 714-16 Point Info ({rsp.number})"]
            for row, zone in zip(range(4, 4 + len(rsp.zones)), rsp.zones):
                target = info.cell(row, 1).value.split("!", 1)[1]
                assert wb["Master"][target].value == f"Z{zone}"
    finally:
        wb.close()
    reread = parse_dmp_worksheet(output)
    assert reread.rsps == design.rsps
    assert reread.keypads == design.keypads
    assert {s.id for s in reread.splitters} == {s.id for s in design.splitters}
    assert {p.number for p in reread.power_supplies} == {p.number for p in design.power_supplies}
