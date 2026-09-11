"""Full LX address bands must remain usable on all door-chart presentation tabs."""

from copy import copy
from pathlib import Path
import re
import sys

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from inject_door_chart import inject, rsp_block_anchor, zone_to_master_row
from parse_dmp_worksheet import DMPDesign, RSP, Zone

TEMPLATE = ROOT / "door_chart_template_blank.xlsx"


def test_lx_edge_mapping_keeps_every_legacy_row_and_block_anchor():
    assert all(zone_to_master_row(number) == number - 434 for number in range(501, 997))
    assert [zone_to_master_row(number) for number in (997, 998, 999, 500)] == [563, 564, 565, 566]
    assert zone_to_master_row(499) == zone_to_master_row(1000) == 0
    assert [rsp_block_anchor(number) for number in (1, 2, 7, 15)] == [67, 83, 167, 299]


@pytest.mark.skipif(not TEMPLATE.exists(), reason="Door-chart template unavailable")
def test_all_bus_edge_addresses_resolve_in_master_and_presentation_tabs(tmp_path):
    design = DMPDesign()
    for number, start in enumerate(
        (start for bus in range(500, 901, 100) for start in (bus, bus + 92)), 1
    ):
        rsp = RSP(number, f"BUILDING {number}", list(range(start, start + 8)), "714-8")
        design.rsps.append(rsp)
        for index, zone in enumerate(rsp.zones):
            design.master_zones.append(Zone(
                zone,
                f"PS-{number}: A/C LOSS" if index == 6 else
                f"PS-{number}: BATT. TRBL" if index == 7 else f"ROOM {zone}",
                number,
                is_ps_ac=index == 6,
                is_ps_batt=index == 7,
            ))
    output = tmp_path / "full-bus-edge-door-chart.xlsx"
    inject(TEMPLATE, design, output)
    wb = openpyxl.load_workbook(output)
    template = openpyxl.load_workbook(TEMPLATE)
    try:
        master = wb["Master"]
        assert master["A66"].value == "ZONE #", "Zone500 must not overwrite the zone-table header"
        for zone in design.master_zones:
            row = zone_to_master_row(zone.number)
            assert master.cell(row, 1).value == f"Z{zone.number}"
            assert master.cell(row, 2).value == zone.description
            if zone.number in (500, 997, 998, 999):
                for column in range(1, 5):
                    cell, original = master.cell(row, column), template["Master"].cell(562, column)
                    assert cell.font == copy(original.font)
                    assert cell.alignment == copy(original.alignment)
                    assert cell.border == copy(original.border)
        for rsp in design.rsps:
            row = zone_to_master_row(rsp.zones[0])
            assert f"BUILDING {rsp.number}" in master.cell(row, 3).value
            assert f"BUILDING {rsp.number}" in master.cell(row, 4).value

        def referenced_zones(sheet_name):
            refs = set()
            for row in wb[sheet_name].iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    assert not re.search(r"Master![A-Z]+0\b", cell.value), "Invalid row-zero formula"
                    for reference in re.findall(r"Master!(A\d+)\b", cell.value):
                        label = master[reference].value
                        if isinstance(label, str) and re.fullmatch(r"Z\d+", label):
                            refs.add(int(label[1:]))
            return refs

        expected = {zone.number for zone in design.master_zones}
        assert referenced_zones("Terminal Cans") == expected
        assert referenced_zones("RSPs") == expected
        assert referenced_zones("Power Supplies") == {
            zone.number for zone in design.master_zones if zone.is_ps_ac or zone.is_ps_batt
        }
    finally:
        template.close()
        wb.close()
