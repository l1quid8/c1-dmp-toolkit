"""Checkpoint printing must compose with riser inventory, not its old stack layout."""

from copy import deepcopy
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.worksheet.cell_range import CellRange
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_dmp_ws import write_dmp_xlsx
from inject_door_chart import inject
from parse_dmp_worksheet import (
    DMPDesign, PowerSupply, RSP, Splitter, Zone, ZoneInfo, parse_dmp_worksheet,
)
from riser_model import DevicePortRef
from topology_service import connect, ensure_explicit_topology
from test_door_chart_consolidation import _rel_targets

CHART_TEMPLATE = ROOT / "door_chart_template_blank.xlsx"
WORKSHEET_TEMPLATE = ROOT / "DMP Installation Worksheet_template_blank.xlsx"
pytestmark = pytest.mark.skipif(not CHART_TEMPLATE.exists(), reason="Chart template unavailable")


def _design(spec):
    """Synthetic inventory; spec is (RSP number, first zone, point count)."""
    design = DMPDesign()
    design.site_info.school_name = "RELEASE CHART TEST"
    for number, first, count in spec:
        zones = list(range(first, first + count))
        location = f"TEST ROOM {number}"
        design.rsps.append(RSP(number, location, zones,
                               "714-8" if count == 8 else "714-16"))
        design.power_supplies.append(PowerSupply(number, location))
        for offset, zone in enumerate(zones):
            ac, batt = offset == count - 2, offset == count - 1
            description = (f"PS-{number}: A/C LOSS" if ac else
                           f"PS-{number}: BATT. TRBL" if batt else f"TEST ZONE {zone}")
            design.master_zones.append(Zone(zone, description, number,
                                            is_ps_ac=ac, is_ps_batt=batt))
            design.zones.append(ZoneInfo(zone, description,
                                         "Supervisory" if ac or batt else "Motion"))
    return design


def _sparse_design():
    # Deliberately not numeric order. Both left and right print columns get an
    # eight-port chart; supply IDs must follow hardware, not physical slot IDs.
    return _design([(15, 717, 16), (2, 501, 8), (11, 701, 8),
                    (4, 509, 16), (7, 601, 16)])


def _assert_sparse_charts(wb):
    terminal, master = wb["Terminal Cans"], wb["Master"]
    for header, first_row, aux, number in (
        ("B7", 67, "B17", 2), ("B32", 75, "B50", 4),
        ("F7", 167, "F25", 7), ("F32", 267, "F42", 11),
        ("B57", 283, "B75", 15),
    ):
        assert terminal[header].value == f"=Master!D{first_row}"
        cell = terminal[aux]
        assert [terminal.cell(cell.row, cell.column + i).value for i in range(3)] == [
            "AUX POWER", f"PS{number}", f"POWER FROM POWER SUPPLY {number}",
        ]
    for col, last, aux, end, zone in ((2, 16, 17, 25, "Z508"),
                                     (6, 41, 42, 50, "Z708")):
        ref = re.search(r"Master!(A\d+)\b", terminal.cell(last, col).value)
        assert ref and master[ref.group(1)].value == zone
        assert all(terminal.cell(row, c).value is None
                   for row in range(aux + 1, end + 1) for c in range(col, col + 3))
    ps_refs = {cell.value for row in wb["Power Supplies"] for cell in row
               if isinstance(cell.value, str)}
    # Actual last two points of RSP 2 and RSP 11, not a 16-point nominal slot.
    for row in (73, 74, 273, 274):
        assert any(re.search(rf"Master!A{row}(?=\D|$)", ref) for ref in ps_refs)
    assert terminal.print_area == (
        "'Terminal Cans'!$B$2:$H$51,'Terminal Cans'!$B$52:$D$76"
    )
    assert terminal.page_setup.scale == 75


def test_sparse_mixed_modules_keep_actual_aux_ids_in_print_order(tmp_path):
    design = _sparse_design()
    before = deepcopy(design)
    output = tmp_path / "sparse-chart.xlsx"
    inject(CHART_TEMPLATE, design, output)
    wb = openpyxl.load_workbook(output)
    try:
        _assert_sparse_charts(wb)
        assert design == before
    finally:
        wb.close()


@pytest.mark.skipif(not WORKSHEET_TEMPLATE.exists(), reason="Worksheet template unavailable")
def test_riser_inventory_roundtrip_keeps_chart_points_and_projected_connections(tmp_path):
    design = _sparse_design()
    design.splitters = [Splitter(sid, family, "TEST SPLITTER ROOM") for sid, family in (
        ("710-LX500-10", "LX"), ("710-KP-2", "KP"), ("710-LX500-2", "LX"),
    )]
    # Initialize schema-2 state as the editor does, then create canonical wiring.
    # Export must project it without rebuilding it from the still-empty legacy fields.
    ensure_explicit_topology(design)
    for src, port, dst in (
        ("MSP", "LX500", "710-LX500-10"),
        ("710-LX500-10", "OUT2", "710-LX500-2"),
        ("710-LX500-2", "OUT1", "RSP-2"),
        ("710-LX500-10", "OUT1", "RSP-4"),
        ("MSP", "PROG", "710-KP-2"),
    ):
        connect(design, DevicePortRef(src, port), DevicePortRef(dst, "IN"))
    connections = deepcopy(design.connections)
    worksheet = tmp_path / "inventory.xlsx"
    write_dmp_xlsx(design, WORKSHEET_TEMPLATE, worksheet)
    assert design.connections == connections
    inventory = openpyxl.load_workbook(worksheet)
    try:
        for sheet, end in (("DMP 714 Exp Mod", 8),
                           ("DMP 505-12_G Power Supply 1-10", 21),
                           ("710 Splitter-Repeater LX500", 9),
                           ("710 Splitter-Repeater(KP-Bus) ", 5)):
            ws = inventory[sheet]
            assert ws.max_row == end
            assert ws.print_area.rsplit("!", 1)[-1] == f"$A$1:$D${end}"
    finally:
        inventory.close()

    reread = parse_dmp_worksheet(worksheet)
    assert reread.rsps == design.rsps
    assert [ps.number for ps in reread.power_supplies] == [15, 2, 11, 4, 7]
    output = tmp_path / "roundtrip-chart.xlsx"
    inject(CHART_TEMPLATE, reread, output)
    wb = openpyxl.load_workbook(output)
    try:
        _assert_sparse_charts(wb)
        master = wb["Master"]
        assert {master.cell(row, 1).value for row in range(67, 563)
                if master.cell(row, 1).value} == {
            f"Z{zone}" for rsp in design.rsps for zone in rsp.zones
        }
        assert [master[f"A{row}"].value for row in (29, 30, 31)] == [
            "710-LX500-2", "710-LX500-10", "710-KP-2",
        ]
        assert master["D29"].value == "From 710-LX500-10"
        assert master["E29"].value == "RSP-2"
        assert master["D30"].value == "500 BUS IN FROM XR/550"
        assert master["F30"].value == "To 710-LX500-2"
        assert [wb["LX-KP-710s"][ref].value for ref in ("B7", "E7", "B19")] == [
            "=Master!C29", "=Master!C30", "=Master!C31",
        ]
    finally:
        wb.close()


@pytest.mark.parametrize("count,last_row,areas", [
    (0, 26, ["$B$2:$D$26"]),
    (1, 26, ["$B$2:$D$26"]),
    (2, 51, ["$B$2:$D$51"]),
    (3, 51, ["$B$2:$H$51"]),
    (4, 51, ["$B$2:$H$51"]),
    (5, 76, ["$B$2:$H$51", "$B$52:$D$76"]),
    (6, 101, ["$B$2:$H$51", "$B$52:$D$101"]),
    (7, 101, ["$B$2:$H$101"]),
    # The template ends at populated row 375; row 376 is print-only spacing.
    (29, 375, ["$B$2:$H$351", "$B$352:$D$376"]),
    (30, 375, ["$B$2:$H$376"]),
])
def test_print_areas_cover_each_chart_once_including_final_template_slot(
        tmp_path, count, last_row, areas):
    design = _design((n, 501 + ((n - 1) // 6) * 100 + ((n - 1) % 6) * 16,
                      8 if n % 2 else 16) for n in range(1, count + 1))
    output = tmp_path / "boundary-chart.xlsx"
    inject(CHART_TEMPLATE, design, output)
    wb = openpyxl.load_workbook(output)
    try:
        tc = wb["Terminal Cans"]
        assert tc.max_row == last_row
        assert tc.print_area == ",".join(f"'Terminal Cans'!{area}" for area in areas)
        assert tc.page_setup.scale == 75
        assert tc.page_setup.pageOrder == "overThenDown"
        assert tc.sheet_properties.pageSetUpPr.fitToPage is False
        assert [brk.id for brk in tc.col_breaks.brk] == [4]
        headers = [cell for row in tc for cell in row
                   if isinstance(cell.value, str) and re.fullmatch(r"=Master!D\d+", cell.value)]
        assert len(headers) == count
        assert len({cell.value for cell in headers}) == count
        printed_ranges = [CellRange(section.rsplit("!", 1)[-1])
                          for section in tc.print_area.split(",")]
        for cell in headers:
            assert sum(cell.coordinate in area for area in printed_ranges) == 1
        if count == 30:
            # The last pair cannot stack in a 15-group template: both physical
            # slots must survive, with the matching real supply labels.
            assert tc["B357"].value == "=Master!D531"  # Z965, RSP 29
            assert tc["F357"].value == "=Master!D547"  # Z981, RSP 30
            assert tc["C367"].value == "PS29"
            assert tc["G375"].value == "PS30"
    finally:
        wb.close()


def test_print_pagination_does_not_relocate_template_tables_or_logos(tmp_path):
    output = tmp_path / "geometry-chart.xlsx"
    inject(CHART_TEMPLATE, _sparse_design(), output)
    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
    with zipfile.ZipFile(CHART_TEMPLATE) as template, zipfile.ZipFile(output) as generated:
        targets = list(_rel_targets(generated, "xl/worksheets/_rels/sheet3.xml.rels"))
        tables = [target for target in targets if target.startswith("xl/tables/")]
        # This template has tables on both sides of groups 1/2 and only the
        # right side of group 3; pagination must retain those exact parts.
        assert set(tables) == {f"xl/tables/table{n}.xml" for n in range(1, 6)}
        for table in tables:
            assert generated.read(table) == template.read(table)
        drawing = next(target for target in targets if target.startswith("xl/drawings/"))
        original = ET.fromstring(template.read(drawing))
        actual = ET.fromstring(generated.read(drawing))
        expected = [anchor for anchor in original
                    if int(anchor.find("xdr:from/xdr:row", ns).text) < 76]
        assert expected
        assert [ET.tostring(anchor) for anchor in actual] == [
            ET.tostring(anchor) for anchor in expected
        ]
